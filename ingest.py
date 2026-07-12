
"""
ingest.py
---------
Turns a PDF (lecture notes, slides, textbook chapter) into searchable pieces.

The pipeline, in plain terms:
1. Pull raw text out of the PDF.
2. Cut that text into overlapping chunks (so no idea gets cut in half).
3. Ask Gemini to turn each chunk into an "embedding" - a list of numbers that
   captures the chunk's meaning, so we can later find chunks that are similar
   in meaning to a question, not just matching keywords.
4. Store those embeddings in a FAISS index - a fast, local search structure
   for "find me the chunks most similar to this new piece of text."

Nothing here is sent anywhere except the text going to Gemini's embedding API.
"""

import os
import pickle
import numpy as np
import faiss
from pypdf import PdfReader
from docx import Document as DocxDocument
from pptx import Presentation
import google.generativeai as genai

DATA_DIR = "data"
INDEX_PATH = os.path.join(DATA_DIR, "index.faiss")
CHUNKS_PATH = os.path.join(DATA_DIR, "chunks.pkl")

EMBED_MODEL = "models/text-embedding-004"
CHUNK_SIZE = 800       # characters per chunk
CHUNK_OVERLAP = 150    # overlap so sentences at the boundary aren't lost


def extract_pages_from_pdf(file_path: str) -> list[dict]:
    """Reads a PDF and returns [{"text": ..., "page": 1}, {"text": ..., "page": 2}, ...]."""
    reader = PdfReader(file_path)
    pages = []
    for i, page in enumerate(reader.pages, start=1):
        text = page.extract_text() or ""
        if text.strip():
            pages.append({"text": text, "page": i})
    return pages


def extract_pages_from_docx(file_path: str) -> list[dict]:
    """
    Reads a Word doc. Word has no real "page" concept in the file itself,
    so we group every ~40 paragraphs into one logical "section" for citation purposes.
    """
    doc = DocxDocument(file_path)
    paras = [p.text for p in doc.paragraphs if p.text.strip()]
    pages = []
    group_size = 40
    for i in range(0, len(paras), group_size):
        section_text = "\n".join(paras[i:i + group_size])
        if section_text.strip():
            pages.append({"text": section_text, "page": (i // group_size) + 1})
    return pages


def extract_pages_from_pptx(file_path: str) -> list[dict]:
    """Reads a PowerPoint file, one 'page' per slide (this maps naturally, unlike docx)."""
    prs = Presentation(file_path)
    pages = []
    for i, slide in enumerate(prs.slides, start=1):
        texts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                texts.append(shape.text_frame.text)
        slide_text = "\n".join(t for t in texts if t.strip())
        if slide_text.strip():
            pages.append({"text": slide_text, "page": i})
    return pages


def extract_pages_from_txt(file_path: str) -> list[dict]:
    """Reads a plain text file as one single 'page' (no natural page concept)."""
    with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
        text = f.read()
    return [{"text": text, "page": 1}] if text.strip() else []


def extract_pages(file_path: str, file_type: str) -> list[dict]:
    """Routes to the right extractor based on file extension."""
    file_type = file_type.lower()
    if file_type == "pdf":
        return extract_pages_from_pdf(file_path)
    elif file_type == "docx":
        return extract_pages_from_docx(file_path)
    elif file_type == "pptx":
        return extract_pages_from_pptx(file_path)
    elif file_type == "txt":
        return extract_pages_from_txt(file_path)
    else:
        raise ValueError(f"Unsupported file type: {file_type}")


def chunk_pages(pages: list[dict], source_name: str, chunk_size=CHUNK_SIZE, overlap=CHUNK_OVERLAP):
    """
    Splits each page's text into overlapping chunks, keeping track of which
    page/slide each chunk came from - this is what powers "Source: notes.pdf, p.4"
    style citations instead of just naming the whole document.
    """
    chunks = []
    for page_data in pages:
        text = page_data["text"].strip()
        page_num = page_data["page"]
        start = 0
        while start < len(text):
            end = start + chunk_size
            chunk = text[start:end].strip()
            if chunk:
                chunks.append({"text": chunk, "source": source_name, "page": page_num})
            start += chunk_size - overlap
    return chunks


def embed_texts(texts: list[str]) -> np.ndarray:
    """
    Calls Gemini's embedding model on a list of strings.
    Returns a numpy array of shape (num_texts, embedding_dim).
    """
    vectors = []
    for t in texts:
        result = genai.embed_content(model=EMBED_MODEL, content=t, task_type="retrieval_document")
        vectors.append(result["embedding"])
    return np.array(vectors, dtype="float32")


def build_or_update_index(new_chunks: list[dict]):
    """
    Embeds new_chunks and adds them to the FAISS index on disk.
    If an index already exists, this appends to it instead of overwriting -
    so uploading a second PDF doesn't erase the first one.
    """
    os.makedirs(DATA_DIR, exist_ok=True)

    texts = [c["text"] for c in new_chunks]
    new_vectors = embed_texts(texts)

    if os.path.exists(INDEX_PATH) and os.path.exists(CHUNKS_PATH):
        index = faiss.read_index(INDEX_PATH)
        with open(CHUNKS_PATH, "rb") as f:
            all_chunks = pickle.load(f)
    else:
        dim = new_vectors.shape[1]
        index = faiss.IndexFlatL2(dim)
        all_chunks = []

    index.add(new_vectors)
    all_chunks.extend(new_chunks)

    faiss.write_index(index, INDEX_PATH)
    with open(CHUNKS_PATH, "wb") as f:
        pickle.dump(all_chunks, f)

    return len(new_chunks)


def load_index():
    """Loads the saved FAISS index and chunk metadata, or returns (None, []) if nothing exists yet."""
    if not (os.path.exists(INDEX_PATH) and os.path.exists(CHUNKS_PATH)):
        return None, []
    index = faiss.read_index(INDEX_PATH)
    with open(CHUNKS_PATH, "rb") as f:
        chunks = pickle.load(f)
    return index, chunks


def list_sources() -> list[str]:
    """Returns the unique document names that have been ingested so far."""
    _, chunks = load_index()
    return sorted(set(c["source"] for c in chunks))


def process_document(file_path: str, source_name: str, file_type: str) -> int:
    """
    Full pipeline for one document (PDF, DOCX, PPTX, or TXT):
    extract -> chunk (with page tracking) -> embed -> store.
    Returns the number of chunks added.
    """
    pages = extract_pages(file_path, file_type)
    chunks = chunk_pages(pages, source_name)
    if not chunks:
        return 0
    return build_or_update_index(chunks)
